import io

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from backend.app.main import app
from backend.ingestion.models import DocumentFileType
from backend.ingestion.service import ingest_document_bytes, validate_file_upload

client = TestClient(app)


def create_sample_pdf_bytes() -> bytes:
    """Generate a 2-page sample PDF in memory using PyMuPDF."""
    doc = pymupdf.open()

    # Page 1
    p1 = doc.new_page()
    p1.insert_text((50, 72), "CONFIDENTIALITY AGREEMENT\n\n1. DEFINITIONS.\nProprietary information is protected.", fontsize=11)

    # Page 2
    p2 = doc.new_page()
    p2.insert_text((50, 72), "2. TERM AND TERMINATION.\nThis agreement shall remain in effect for 3 years.", fontsize=11)

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes() -> bytes:
    """Generate a sample DOCX with paragraphs and tables in memory."""
    doc = docx.Document()
    doc.add_heading("NON-DISCLOSURE AGREEMENT", level=1)
    doc.add_paragraph("This Agreement is entered into between Party A and Party B.")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Party"
    table.rows[0].cells[1].text = "Role"
    table.rows[1].cells[0].text = "Acme Corp"
    table.rows[1].cells[1].text = "Discloser"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_parsing_and_offset_continuity():
    pdf_bytes = create_sample_pdf_bytes()
    parsed = ingest_document_bytes(pdf_bytes, "sample_nda.pdf")

    assert parsed.file_type == DocumentFileType.PDF
    assert parsed.page_count == 2
    assert len(parsed.pages) == 2
    assert "CONFIDENTIALITY AGREEMENT" in parsed.full_text
    assert "TERM AND TERMINATION" in parsed.full_text

    # Invariant: Each span char_start_doc to char_end_doc must match span text
    for page in parsed.pages:
        for span in page.spans:
            extracted_from_doc = parsed.full_text[span.char_start_doc:span.char_end_doc]
            assert extracted_from_doc == span.text, f"Offset mismatch: '{extracted_from_doc}' != '{span.text}'"


def test_docx_parsing_and_tables():
    docx_bytes = create_sample_docx_bytes()
    parsed = ingest_document_bytes(docx_bytes, "sample_contract.docx")

    assert parsed.file_type == DocumentFileType.DOCX
    assert parsed.page_count >= 1
    assert "NON-DISCLOSURE AGREEMENT" in parsed.full_text
    assert "Acme Corp" in parsed.full_text
    assert "Discloser" in parsed.full_text


def test_invalid_extension_rejection():
    with pytest.raises(Exception) as exc:
        validate_file_upload("contract.exe", 1024)
    assert "Unsupported file type" in str(exc.value)


def test_empty_file_rejection():
    with pytest.raises(Exception) as exc:
        validate_file_upload("empty.pdf", 0)
    assert "empty (0 bytes)" in str(exc.value)


def test_upload_api_roundtrip():
    pdf_bytes = create_sample_pdf_bytes()
    files = {"file": ("test_api_doc.pdf", pdf_bytes, "application/pdf")}

    # 1. Upload
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 201
    doc_data = response.json()
    doc_id = doc_data["id"]
    assert doc_data["filename"] == "test_api_doc.pdf"
    assert doc_data["page_count"] == 2

    # 2. Get document metadata
    get_res = client.get(f"/api/documents/{doc_id}")
    assert get_res.status_code == 200
    assert get_res.json()["id"] == doc_id

    # 3. Get page 1
    p1_res = client.get(f"/api/documents/{doc_id}/pages/1")
    assert p1_res.status_code == 200
    assert p1_res.json()["page_number"] == 1

    # 4. Get full text
    text_res = client.get(f"/api/documents/{doc_id}/text")
    assert text_res.status_code == 200
    assert "CONFIDENTIALITY AGREEMENT" in text_res.json()["text"]

    # 5. List documents
    list_res = client.get("/api/documents")
    assert list_res.status_code == 200
    assert any(d["id"] == doc_id for d in list_res.json())

    # 6. Delete document
    del_res = client.delete(f"/api/documents/{doc_id}")
    assert del_res.status_code == 200
